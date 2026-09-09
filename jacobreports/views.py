from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q, Sum
from django.http import HttpResponse
from datetime import date, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from institutions.models import school
from setup.models import academicyr,currentacademicyr,sclass,section
from staff.models import staff
from .models import ClassTeacherReport, InchargeReport, ClassTeacherMapping, InchargeMapping, PrincipalDailyLog, PrincipalLogAccessMapping, PrincipalLogEntryType
from .forms import ClassTeacherReportForm, InchargeReportForm, ClassTeacherMappingForm, InchargeMappingForm, PrincipalLogEntryForm, PrincipalLogAccessForm


# =========================================
# REPORT LIST
# =========================================

def class_teacher_report_list(request):
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)
    usr = request.user

    reports = ClassTeacherReport.objects.select_related(
        'school_name',
        'class_name',
        'section',
        'report_submitted_by'
    ).order_by('-report_date', '-id')

    is_admin = usr.groups.filter(name__in=['superadmin', 'Admin']).exists()

    if is_admin:
        reports = reports.filter(school_name=sdata)
    else:
        try:
            current_staff = staff.objects.get(staff_user=usr)
        except staff.DoesNotExist:
            reports = reports.none()
        else:
            mapped_teacher_ids = ClassTeacherMapping.objects.filter(
                incharge=current_staff,
                school_name=sdata
            ).values_list('class_teacher', flat=True)

            # always see own report + all mapped teachers' reports
            reports = reports.filter(
                school_name=sdata
            ).filter(
                Q(report_submitted_by=current_staff) |
                Q(report_submitted_by__in=mapped_teacher_ids)
            )

    # =========================
    # SEARCH
    # =========================

    search = request.GET.get('search')

    if search:
        reports = reports.filter(
            Q(class_name__class_name__icontains=search) |
            Q(section__section__icontains=search) |
            Q(report_submitted_by__first_name__icontains=search) |
            Q(report_submitted_by__last_name__icontains=search)
        )

    # =========================
    # DATE FILTER
    # =========================

    report_date = request.GET.get('report_date')

    if report_date:
        reports = reports.filter(report_date=report_date)

    # =========================
    # PAGINATION
    # =========================

    paginator = Paginator(reports, 20)

    page_number = request.GET.get('page')

    page_obj = paginator.get_page(page_number)

    context = {
        'page_obj': page_obj,
        'search': search,
        'report_date': report_date,
        'sdata': sdata,
        'is_admin': is_admin,
    }

    return render(
        request,
        'jacobreports/class_teacher_report_list.html',
        context
    )


# =========================================
# CONSOLIDATED REPORT (ALL CLASSES, ONE DATE)
# =========================================

def consolidated_class_teacher_report(request):
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    report_date = request.GET.get('report_date') or date.today().isoformat()

    try:
        current_date_obj = date.fromisoformat(report_date)
    except ValueError:
        current_date_obj = date.today()
        report_date = current_date_obj.isoformat()

    prev_date = (current_date_obj - timedelta(days=1)).isoformat()
    next_date = (current_date_obj + timedelta(days=1)).isoformat()

    usr = request.user

    reports = ClassTeacherReport.objects.select_related(
        'class_name',
        'section',
        'report_submitted_by'
    ).filter(
        school_name=sdata,
        report_date=report_date
    ).order_by('class_name__name', 'section__section_name')

    is_admin = usr.groups.filter(name__in=['superadmin', 'Admin']).exists()

    if not is_admin:
        try:
            current_staff = staff.objects.get(staff_user=usr)
        except staff.DoesNotExist:
            reports = reports.none()
        else:
            mapped_teacher_ids = ClassTeacherMapping.objects.filter(
                incharge=current_staff,
                school_name=sdata
            ).values_list('class_teacher', flat=True)

            reports = reports.filter(
                Q(report_submitted_by=current_staff) |
                Q(report_submitted_by__in=mapped_teacher_ids)
            )

    # =========================
    # SCHOOL-WIDE TOTALS
    # =========================

    sums = reports.aggregate(
        boys_on_roll=Sum('boys_on_roll'),
        girls_on_roll=Sum('girls_on_roll'),
        boys_present=Sum('boys_present'),
        girls_present=Sum('girls_present'),
        boys_absentees=Sum('boys_absentees'),
        girls_absentees=Sum('girls_absentees'),
        boys_uniform_defaulters=Sum('boys_uniform_defaulters'),
        girls_uniform_defaulters=Sum('girls_uniform_defaulters'),
        private_auto_boys=Sum('private_auto_boys'),
        private_auto_girls=Sum('private_auto_girls'),
        cycle_boys=Sum('cycle_boys'),
        cycle_girls=Sum('cycle_girls'),
        walk_boys=Sum('walk_boys'),
        walk_girls=Sum('walk_girls'),
        school_van_boys=Sum('school_van_boys'),
        school_van_girls=Sum('school_van_girls'),
        bus_boys=Sum('bus_boys'),
        bus_girls=Sum('bus_girls'),
        others_boys=Sum('others_boys'),
        others_girls=Sum('others_girls'),
    )

    for key, value in sums.items():
        sums[key] = value or 0

    total_on_roll = sums['boys_on_roll'] + sums['girls_on_roll']
    total_present = sums['boys_present'] + sums['girls_present']
    total_absentees = sums['boys_absentees'] + sums['girls_absentees']
    total_uniform_defaulters = (
        sums['boys_uniform_defaulters'] +
        sums['girls_uniform_defaulters']
    )

    attendance_pct = round(
        (total_present / total_on_roll) * 100, 1
    ) if total_on_roll else 0

    transport_rows = [
        {
            'mode': 'Private Auto',
            'boys': sums['private_auto_boys'],
            'girls': sums['private_auto_girls'],
            'total': sums['private_auto_boys'] + sums['private_auto_girls'],
        },
        {
            'mode': 'Cycle',
            'boys': sums['cycle_boys'],
            'girls': sums['cycle_girls'],
            'total': sums['cycle_boys'] + sums['cycle_girls'],
        },
        {
            'mode': 'Walk',
            'boys': sums['walk_boys'],
            'girls': sums['walk_girls'],
            'total': sums['walk_boys'] + sums['walk_girls'],
        },
        {
            'mode': 'School Van',
            'boys': sums['school_van_boys'],
            'girls': sums['school_van_girls'],
            'total': sums['school_van_boys'] + sums['school_van_girls'],
        },
        {
            'mode': 'Bus',
            'boys': sums['bus_boys'],
            'girls': sums['bus_girls'],
            'total': sums['bus_boys'] + sums['bus_girls'],
        },
        {
            'mode': 'Others',
            'boys': sums['others_boys'],
            'girls': sums['others_girls'],
            'total': sums['others_boys'] + sums['others_girls'],
        },
    ]

    transport_grand_total = sum(row['total'] for row in transport_rows)

    # =========================
    # PENDING (NOT YET SUBMITTED) SECTIONS
    # =========================

    all_sections = section.objects.none()
    pending_sections = []

    if is_admin:
        try:
            cur_yr = currentacademicyr.objects.get(school_name=sdata)
            all_sections = section.objects.filter(
                school_name=sdata,
                acad_year=cur_yr
            ).select_related('class_sec_name').order_by(
                'class_sec_name__name',
                'section_name'
            )
            submitted_pairs = set(
                reports.values_list('class_name_id', 'section_id')
            )
            pending_sections = [
                sec for sec in all_sections
                if (sec.class_sec_name_id, sec.id) not in submitted_pairs
            ]
        except currentacademicyr.DoesNotExist:
            pass

    # =========================
    # ITEMS NEEDING ATTENTION
    # =========================

    def _has_content(value):
        return bool(value) and value.strip().lower() not in ['nil', 'none', 'na', 'n/a']

    flagged_reports = [
        r for r in reports
        if _has_content(r.accident_details)
        or _has_content(r.defaulters)
        or r.total_absentees > 0
    ]

    _text_fields = [
        'teachers_remark', 'parents_remark', 'pupils_remark',
        'action_taken', 'birthday_celebration', 'homework_details',
        'drill_work_details', 'activity_class', 'announcements',
        'defaulters', 'accident_details', 'meeting_details',
        'suggestions_and_grievances',
    ]

    reports_with_remarks = [
        r for r in reports
        if any(getattr(r, f) and str(getattr(r, f)).strip() for f in _text_fields)
    ]

    context = {
        'report_date': report_date,
        'prev_date': prev_date,
        'next_date': next_date,
        'is_today': report_date == date.today().isoformat(),
        'reports': reports,
        'sums': sums,
        'total_on_roll': total_on_roll,
        'total_present': total_present,
        'total_absentees': total_absentees,
        'total_uniform_defaulters': total_uniform_defaulters,
        'attendance_pct': attendance_pct,
        'classes_reported': reports.count(),
        'classes_total': all_sections.count() if is_admin else reports.count(),
        'transport_rows': transport_rows,
        'transport_grand_total': transport_grand_total,
        'pending_sections': pending_sections,
        'flagged_reports': flagged_reports,
        'reports_with_remarks': reports_with_remarks,
    }

    return render(
        request,
        'jacobreports/class_teacher_report_consolidated.html',
        context
    )


# =========================================
# VIEW REPORT DETAIL
# =========================================

def view_class_teacher_report(request, report_id):
    report = get_object_or_404(
        ClassTeacherReport.objects.select_related(
            'school_name', 'class_name', 'section', 'report_submitted_by'
        ),
        id=report_id
    )
    return render(request, 'jacobreports/class_teacher_report_detail.html', {'report': report})


# =========================================
# ADD REPORT
# =========================================

def add_class_teacher_report(request):
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)
    usr = request.user

    is_admin = usr.groups.filter(name__in=['superadmin', 'Admin']).exists()

    if is_admin:
        staff_qs = staff.objects.filter(staff_school=sdata)
    else:
        staff_qs = staff.objects.filter(staff_user=usr)

    # =========================
    # CLONE PREVIOUS REPORT
    # =========================

    clone_id = request.GET.get('clone')

    initial_data = {

    }

    if clone_id:

        old_report = get_object_or_404(
            ClassTeacherReport,
            id=clone_id
        )

        for field in ClassTeacherReport._meta.fields:

            if field.name not in [
                'id',
                'created_at',
                'report_date'
            ]:
                initial_data[field.name] = getattr(
                    old_report,
                    field.name
                )

        initial_data['report_date'] = date.today()
    elif is_admin:
        first_staff = staff_qs.first()
        if first_staff:
            initial_data['report_submitted_by'] = first_staff.pk

    # =========================
    # SAVE FORM
    # =========================

    if request.method == 'POST':

        form = ClassTeacherReportForm(
            request.POST
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                'Class teacher report added successfully.'
            )

            return redirect(
                'class_teacher_report_list'
            )

    else:

        form = ClassTeacherReportForm(
            initial=initial_data
        )

        form.fields['school_name'].queryset = school.objects.filter(id=sch_id)
        form.fields['class_name'].queryset = sclass.objects.filter(school_name=sdata)
        form.fields['report_submitted_by'].queryset = staff_qs
    context = {
        'form': form
    }

    return render(
        request,
        'jacobreports/class_teacher_report_form.html',
        context
    )


# =========================================
# EDIT REPORT
# =========================================

def edit_class_teacher_report(request, report_id):

    report = get_object_or_404(
        ClassTeacherReport,
        id=report_id
    )

    if request.method == 'POST':

        form = ClassTeacherReportForm(
            request.POST,
            instance=report
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                'Report updated successfully.'
            )

            return redirect(
                'class_teacher_report_list'
            )

    else:

        form = ClassTeacherReportForm(
            instance=report
        )

    context = {
        'form': form,
        'report': report
    }

    return render(
        request,
        'jacobreports/edit_class_teacher_report.html',
        context
    )


# =========================================
# DELETE REPORT
# =========================================

def delete_class_teacher_report(request, report_id):

    report = get_object_or_404(
        ClassTeacherReport,
        id=report_id
    )

    report.delete()

    messages.success(
        request,
        'Report deleted successfully.'
    )

    return redirect(
        'class_teacher_report_list'
    )


# =========================================
# INCHARGE REPORT — LIST
# =========================================

def incharge_report_list(request):
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)
    usr = request.user

    is_admin = usr.groups.filter(name__in=['superadmin', 'Admin', 'Accounts']).exists()

    reports = InchargeReport.objects.filter(
        school_name=sdata
    ).select_related(
        'school_name', 'report_submitted_by'
    )

    if not is_admin:
        try:
            current_staff = staff.objects.get(staff_user=usr)
        except staff.DoesNotExist:
            reports = reports.none()
        else:
            mapped_incharge_ids = InchargeMapping.objects.filter(
                supervisor=current_staff,
                school_name=sdata
            ).values_list('incharge', flat=True)

            reports = reports.filter(
                Q(report_submitted_by=current_staff) |
                Q(report_submitted_by__in=mapped_incharge_ids)
            )

    search = request.GET.get('search', '')
    report_date = request.GET.get('report_date', '')
    department = request.GET.get('department', '')

    if search:
        reports = reports.filter(
            Q(report_submitted_by__first_name__icontains=search) |
            Q(report_submitted_by__last_name__icontains=search) |
            Q(department__icontains=search)
        )

    if report_date:
        reports = reports.filter(report_date=report_date)

    if department:
        reports = reports.filter(department=department)

    paginator = Paginator(reports, 20)
    page_obj = paginator.get_page(request.GET.get('page'))

    context = {
        'page_obj': page_obj,
        'search': search,
        'report_date': report_date,
        'department': department,
        'department_choices': InchargeReport.DEPARTMENT_CHOICES,
        'is_admin': is_admin,
        'sdata': sdata,
    }

    return render(request, 'jacobreports/incharge_report_list.html', context)


# =========================================
# INCHARGE REPORT — DETAIL
# =========================================

def view_incharge_report(request, report_id):
    report = get_object_or_404(
        InchargeReport.objects.select_related('school_name', 'report_submitted_by'),
        id=report_id
    )
    return render(request, 'jacobreports/incharge_report_detail.html', {'report': report})


# =========================================
# INCHARGE REPORT — ADD
# =========================================

def add_incharge_report(request):
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)
    usr = request.user

    is_superadmin = usr.groups.filter(name='superadmin').exists()

    if is_superadmin:
        staff_qs = staff.objects.filter(staff_school=sdata)
    else:
        staff_qs = staff.objects.filter(staff_user=usr)

    clone_id = request.GET.get('clone')
    initial_data = {}

    if clone_id:
        old = get_object_or_404(InchargeReport, id=clone_id)
        for field in InchargeReport._meta.fields:
            if field.name not in ['id', 'created_at', 'report_date']:
                initial_data[field.name] = getattr(old, field.name)
        initial_data['report_date'] = date.today()
    elif is_superadmin:
        first_staff = staff_qs.first()
        if first_staff:
            initial_data['report_submitted_by'] = first_staff.pk

    if request.method == 'POST':
        form = InchargeReportForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Incharge report added successfully.')
            return redirect('incharge_report_list')
    else:
        form = InchargeReportForm(initial=initial_data)
        form.fields['school_name'].queryset = school.objects.filter(id=sch_id)
        form.fields['report_submitted_by'].queryset = staff_qs

    return render(request, 'jacobreports/incharge_report_form.html', {'form': form, 'title': 'ADD INCHARGE REPORT'})


# =========================================
# INCHARGE REPORT — EDIT
# =========================================

def edit_incharge_report(request, report_id):
    report = get_object_or_404(InchargeReport, id=report_id)

    if request.method == 'POST':
        form = InchargeReportForm(request.POST, instance=report)
        if form.is_valid():
            form.save()
            messages.success(request, 'Incharge report updated successfully.')
            return redirect('incharge_report_list')
    else:
        form = InchargeReportForm(instance=report)

    return render(request, 'jacobreports/incharge_report_form.html', {
        'form': form,
        'report': report,
        'title': 'EDIT INCHARGE REPORT'
    })


# =========================================
# INCHARGE REPORT — DELETE
# =========================================

def delete_incharge_report(request, report_id):
    report = get_object_or_404(InchargeReport, id=report_id)
    report.delete()
    messages.success(request, 'Incharge report deleted successfully.')
    return redirect('incharge_report_list')


# =========================================
# CLASS TEACHER MAPPING — MANAGE
# =========================================

def manage_class_teacher_mappings(request):
    usr = request.user
    if not usr.groups.filter(name__in=['superadmin', 'Admin', 'Accounts']).exists():
        messages.error(request, 'You do not have permission to access this page.')
        return redirect('class_teacher_report_list')

    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    staff_qs = staff.objects.filter(staff_school=sdata).order_by('first_name', 'last_name')

    if request.method == 'POST':
        form = ClassTeacherMappingForm(request.POST)
        form.fields['incharge'].queryset = staff_qs
        form.fields['class_teachers'].queryset = staff_qs
        if form.is_valid():
            incharge_staff = form.cleaned_data['incharge']
            selected_teachers = form.cleaned_data['class_teachers']
            added, skipped = 0, 0
            for teacher in selected_teachers:
                _, created = ClassTeacherMapping.objects.get_or_create(
                    school_name=sdata,
                    incharge=incharge_staff,
                    class_teacher=teacher
                )
                if created:
                    added += 1
                else:
                    skipped += 1
            if added:
                messages.success(request, f'{added} mapping(s) added successfully.')
            if skipped:
                messages.warning(request, f'{skipped} mapping(s) already existed and were skipped.')
        else:
            messages.error(request, 'Please correct the errors below.')
        return redirect('manage_class_teacher_mappings')

    form = ClassTeacherMappingForm()
    form.fields['incharge'].queryset = staff_qs
    form.fields['class_teachers'].queryset = staff_qs

    mappings = ClassTeacherMapping.objects.filter(
        school_name=sdata
    ).select_related('incharge', 'class_teacher').order_by(
        'incharge__first_name', 'class_teacher__first_name'
    )

    context = {
        'form': form,
        'mappings': mappings,
    }

    return render(request, 'jacobreports/class_teacher_mappings.html', context)


# =========================================
# CLASS TEACHER MAPPING — DELETE
# =========================================

def delete_class_teacher_mapping(request, mapping_id):
    usr = request.user
    if not usr.groups.filter(name__in=['superadmin', 'Admin', 'Accounts']).exists():
        messages.error(request, 'You do not have permission to perform this action.')
        return redirect('class_teacher_report_list')

    mapping = get_object_or_404(ClassTeacherMapping, id=mapping_id)
    mapping.delete()
    messages.success(request, 'Mapping removed successfully.')
    return redirect('manage_class_teacher_mappings')


# =========================================
# BACKUP TO EXCEL & DELETE
# =========================================

def export_and_delete_class_teacher_reports(request):
    if not request.user.groups.filter(name__in=['superadmin', 'Admin']).exists():
        messages.error(request, 'You do not have permission to perform this action.')
        return redirect('class_teacher_report_list')

    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')

    if not from_date or not to_date:
        messages.error(request, 'Please provide both from date and to date.')
        return redirect('class_teacher_report_list')

    reports = ClassTeacherReport.objects.select_related(
        'school_name', 'class_name', 'section', 'report_submitted_by'
    ).filter(
        school_name=sdata,
        report_date__gte=from_date,
        report_date__lte=to_date,
    ).order_by('report_date', 'class_name__name', 'section__section_name')

    if not reports.exists():
        messages.warning(request, 'No reports found for the selected date range.')
        return redirect('class_teacher_report_list')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Class Teacher Reports'

    header_fill = PatternFill('solid', fgColor='212529')
    header_font = Font(color='FFFFFF', bold=True)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    headers = [
        'Date', 'School', 'Class', 'Section', 'Submitted By',
        'Boys on Roll', 'Girls on Roll', 'Total on Roll',
        'Boys Present', 'Girls Present', 'Total Present',
        'Boys Absent', 'Girls Absent', 'Total Absent',
        'Boys Uniform Defaulters', 'Girls Uniform Defaulters', 'Total Uniform Defaulters',
        'Attendance %',
        'Action Taken', 'Birthday Celebration',
        'Accident Details', 'Defaulters', 'Homework Details',
        'Drill Work Details', 'Activity Class', 'Announcements',
        'Teachers Remark', 'Parents Remark', 'Pupils Remark',
        'Private Auto Boys', 'Private Auto Girls', 'Private Auto Total',
        'Cycle Boys', 'Cycle Girls', 'Cycle Total',
        'Walk Boys', 'Walk Girls', 'Walk Total',
        'School Van Boys', 'School Van Girls', 'School Van Total',
        'Bus Boys', 'Bus Girls', 'Bus Total',
        'Others Boys', 'Others Girls', 'Others Total',
        'Transport Grand Total',
        'Meeting Details', 'Suggestions & Grievances',
        'Created At',
    ]

    ws.append(headers)
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    for r in reports:
        ws.append([
            r.report_date.strftime('%Y-%m-%d'),
            str(r.school_name),
            str(r.class_name),
            str(r.section),
            str(r.report_submitted_by),
            r.boys_on_roll, r.girls_on_roll, r.total_on_roll,
            r.boys_present, r.girls_present, r.total_present,
            r.boys_absentees, r.girls_absentees, r.total_absentees,
            r.boys_uniform_defaulters, r.girls_uniform_defaulters, r.total_uniform_defaulters,
            r.attendance_pct,
            r.action_taken or '', r.birthday_celebration or '',
            r.accident_details or '', r.defaulters or '', r.homework_details or '',
            r.drill_work_details or '', r.activity_class or '', r.announcements or '',
            r.teachers_remark or '', r.parents_remark or '', r.pupils_remark or '',
            r.private_auto_boys, r.private_auto_girls, r.private_auto_total,
            r.cycle_boys, r.cycle_girls, r.cycle_total,
            r.walk_boys, r.walk_girls, r.walk_total,
            r.school_van_boys, r.school_van_girls, r.school_van_total,
            r.bus_boys, r.bus_girls, r.bus_total,
            r.others_boys, r.others_girls, r.others_total,
            r.transport_grand_total,
            r.meeting_details or '', r.suggestions_and_grievances or '',
            r.created_at.strftime('%Y-%m-%d %H:%M'),
        ])

    for col in ws.columns:
        max_len = max((len(str(cell.value)) if cell.value else 0) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    total_exported = reports.count()
    filename = f"class_teacher_reports_{from_date}_to_{to_date}.xlsx"

    # Find the latest date in the range — keep those records in the database
    latest_date = reports.order_by('-report_date').values_list('report_date', flat=True).first()
    to_delete = reports.exclude(report_date=latest_date)
    deleted_count = to_delete.count()
    to_delete.delete()

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)

    messages.success(
        request,
        f'{total_exported} report(s) exported to Excel. '
        f'{deleted_count} record(s) deleted (reports from {latest_date} kept in database).'
    )
    return response


# =========================================
# INCHARGE MAPPING — MANAGE
# =========================================

def manage_incharge_mappings(request):
    usr = request.user
    if not usr.groups.filter(name__in=['superadmin', 'Admin', 'Accounts']).exists():
        messages.error(request, 'You do not have permission to access this page.')
        return redirect('incharge_report_list')

    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    staff_qs = staff.objects.filter(staff_school=sdata).order_by('first_name', 'last_name')

    if request.method == 'POST':
        form = InchargeMappingForm(request.POST)
        form.fields['supervisor'].queryset = staff_qs
        form.fields['incharges'].queryset = staff_qs
        if form.is_valid():
            supervisor_staff = form.cleaned_data['supervisor']
            selected_incharges = form.cleaned_data['incharges']
            added, skipped = 0, 0
            for incharge_staff in selected_incharges:
                _, created = InchargeMapping.objects.get_or_create(
                    school_name=sdata,
                    supervisor=supervisor_staff,
                    incharge=incharge_staff
                )
                if created:
                    added += 1
                else:
                    skipped += 1
            if added:
                messages.success(request, f'{added} mapping(s) added successfully.')
            if skipped:
                messages.warning(request, f'{skipped} mapping(s) already existed and were skipped.')
        else:
            messages.error(request, 'Please correct the errors below.')
        return redirect('manage_incharge_mappings')

    form = InchargeMappingForm()
    form.fields['supervisor'].queryset = staff_qs
    form.fields['incharges'].queryset = staff_qs

    mappings = InchargeMapping.objects.filter(
        school_name=sdata
    ).select_related('supervisor', 'incharge').order_by(
        'supervisor__first_name', 'incharge__first_name'
    )

    context = {
        'form': form,
        'mappings': mappings,
    }

    return render(request, 'jacobreports/incharge_mappings.html', context)


# =========================================
# INCHARGE MAPPING — DELETE
# =========================================

def delete_incharge_mapping(request, mapping_id):
    usr = request.user
    if not usr.groups.filter(name__in=['superadmin', 'Admin', 'Accounts']).exists():
        messages.error(request, 'You do not have permission to perform this action.')
        return redirect('incharge_report_list')

    mapping = get_object_or_404(InchargeMapping, id=mapping_id)
    mapping.delete()
    messages.success(request, 'Mapping removed successfully.')
    return redirect('manage_incharge_mappings')


# =========================================
# PRINCIPAL DAILY LOG — ACCESS HELPER
# =========================================

def _has_principal_log_access(user, sdata):
    """Returns True if user may do CRUD on PrincipalDailyLog."""
    if user.groups.filter(name__in=['superadmin', 'Admin']).exists():
        return True
    try:
        current_staff = staff.objects.get(staff_user=user)
    except staff.DoesNotExist:
        return False
    return PrincipalLogAccessMapping.objects.filter(
        school_name=sdata,
        staff_member=current_staff
    ).exists()


# =========================================
# PRINCIPAL DAILY LOG — LIST
# =========================================

def principal_log_list(request):
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    # Auto-seed default types on first visit
    if not PrincipalLogEntryType.objects.filter(school_name=sdata).exists():
        PrincipalLogEntryType.seed_defaults(sdata)

    entries = PrincipalDailyLog.objects.filter(school_name=sdata)

    entry_date = request.GET.get('entry_date', '')
    entry_type = request.GET.get('entry_type', '')
    status = request.GET.get('status', '')
    search = request.GET.get('search', '')

    if entry_date:
        entries = entries.filter(entry_date=entry_date)
    if entry_type:
        entries = entries.filter(entry_type=entry_type)
    if status:
        entries = entries.filter(status=status)
    if search:
        entries = entries.filter(description__icontains=search)

    paginator = Paginator(entries, 25)
    page_obj = paginator.get_page(request.GET.get('page'))

    is_admin = request.user.groups.filter(name__in=['superadmin', 'Admin']).exists()
    can_crud = _has_principal_log_access(request.user, sdata)

    entry_types = PrincipalLogEntryType.objects.filter(school_name=sdata)

    context = {
        'page_obj': page_obj,
        'sdata': sdata,
        'entry_date': entry_date,
        'entry_type': entry_type,
        'status': status,
        'search': search,
        'entry_types': entry_types,
        'status_choices': PrincipalDailyLog.STATUS_CHOICES,
        'can_crud': can_crud,
        'is_admin': is_admin,
    }
    return render(request, 'jacobreports/principal_log_list.html', context)


# =========================================
# PRINCIPAL DAILY LOG — ADD
# =========================================

def add_principal_log_entry(request):
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    if not _has_principal_log_access(request.user, sdata):
        messages.error(request, 'You do not have permission to add entries.')
        return redirect('principal_log_list')

    if request.method == 'POST':
        form = PrincipalLogEntryForm(request.POST, school=sdata)
        if form.is_valid():
            form.save()
            messages.success(request, 'Entry added successfully.')
            return redirect('principal_log_list')
    else:
        form = PrincipalLogEntryForm(school=sdata, initial={
            'school_name': sdata.pk,
            'entry_date': date.today(),
        })
        form.fields['school_name'].queryset = school.objects.filter(id=sch_id)

    return render(request, 'jacobreports/principal_log_form.html', {
        'form': form,
        'title': 'ADD ENTRY',
        'sdata': sdata,
        'can_crud': _has_principal_log_access(request.user, sdata),
    })


# =========================================
# PRINCIPAL DAILY LOG — EDIT
# =========================================

def edit_principal_log_entry(request, entry_id):
    entry = get_object_or_404(PrincipalDailyLog, id=entry_id)
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    if not _has_principal_log_access(request.user, sdata):
        messages.error(request, 'You do not have permission to edit entries.')
        return redirect('principal_log_list')

    if request.method == 'POST':
        form = PrincipalLogEntryForm(request.POST, instance=entry, school=sdata)
        if form.is_valid():
            form.save()
            messages.success(request, 'Entry updated successfully.')
            return redirect('principal_log_list')
    else:
        form = PrincipalLogEntryForm(instance=entry, school=sdata)
        form.fields['school_name'].queryset = school.objects.filter(id=sch_id)

    return render(request, 'jacobreports/principal_log_form.html', {
        'form': form,
        'title': 'EDIT ENTRY',
        'entry': entry,
        'sdata': sdata,
        'can_crud': _has_principal_log_access(request.user, sdata),
    })


# =========================================
# PRINCIPAL DAILY LOG — DELETE
# =========================================

def delete_principal_log_entry(request, entry_id):
    entry = get_object_or_404(PrincipalDailyLog, id=entry_id)
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    if not _has_principal_log_access(request.user, sdata):
        messages.error(request, 'You do not have permission to delete entries.')
        return redirect('principal_log_list')

    entry.delete()
    messages.success(request, 'Entry deleted successfully.')
    return redirect('principal_log_list')


# =========================================
# PRINCIPAL DAILY LOG — PRINT REPORT
# =========================================

def principal_log_report(request):
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    from_date = request.GET.get('from_date', '')
    to_date = request.GET.get('to_date', '')
    status_filter = request.GET.get('status', '')

    entries = PrincipalDailyLog.objects.filter(school_name=sdata)

    if from_date:
        entries = entries.filter(entry_date__gte=from_date)
    if to_date:
        entries = entries.filter(entry_date__lte=to_date)
    if status_filter:
        entries = entries.filter(status=status_filter)

    entries = entries.order_by('-entry_date', 'entry_type', 'id')

    # Group: date → type → entries
    report_data = []
    dates_seen = []
    date_map = {}
    for entry in entries:
        d = entry.entry_date
        t = entry.entry_type
        if d not in date_map:
            date_map[d] = {}
            dates_seen.append(d)
        if t not in date_map[d]:
            date_map[d][t] = []
        date_map[d][t].append(entry)

    for d in dates_seen:
        type_groups = [
            {'type': t, 'entries': date_map[d][t]}
            for t in date_map[d]
        ]
        report_data.append({'date': d, 'type_groups': type_groups})

    context = {
        'report_data': report_data,
        'sdata': sdata,
        'from_date': from_date,
        'to_date': to_date,
        'status_filter': status_filter,
        'status_choices': PrincipalDailyLog.STATUS_CHOICES,
        'total_entries': entries.count(),
        'entry_types': PrincipalLogEntryType.objects.filter(school_name=sdata),
    }
    return render(request, 'jacobreports/principal_log_report.html', context)


# =========================================
# PRINCIPAL LOG ENTRY TYPES — MANAGE PAGE
# =========================================

def manage_principal_log_types(request):
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    if not _has_principal_log_access(request.user, sdata):
        messages.error(request, 'Permission denied.')
        return redirect('principal_log_list')

    if not PrincipalLogEntryType.objects.filter(school_name=sdata).exists():
        PrincipalLogEntryType.seed_defaults(sdata)

    if request.method == 'POST':
        name = request.POST.get('type_name', '').strip()
        if name:
            _, created = PrincipalLogEntryType.objects.get_or_create(
                school_name=sdata, name=name
            )
            if created:
                messages.success(request, f'Type "{name}" added.')
            else:
                messages.warning(request, f'Type "{name}" already exists.')
        else:
            messages.error(request, 'Type name cannot be empty.')
        return redirect('manage_principal_log_types')

    entry_types = PrincipalLogEntryType.objects.filter(school_name=sdata)

    return render(request, 'jacobreports/principal_log_types.html', {
        'sdata': sdata,
        'entry_types': entry_types,
    })


# =========================================
# PRINCIPAL LOG ENTRY TYPES — ADD (inline redirect)
# =========================================

def add_principal_log_type(request):
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    if not _has_principal_log_access(request.user, sdata):
        messages.error(request, 'Permission denied.')
        return redirect('principal_log_list')

    if request.method == 'POST':
        name = request.POST.get('type_name', '').strip()
        if name:
            _, created = PrincipalLogEntryType.objects.get_or_create(
                school_name=sdata, name=name
            )
            if created:
                messages.success(request, f'Type "{name}" added.')
            else:
                messages.warning(request, f'Type "{name}" already exists.')
        else:
            messages.error(request, 'Type name cannot be empty.')

    return redirect('principal_log_list')


# =========================================
# PRINCIPAL LOG ENTRY TYPES — DELETE
# =========================================

def delete_principal_log_type(request, type_id):
    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    if not _has_principal_log_access(request.user, sdata):
        messages.error(request, 'Permission denied.')
        return redirect('principal_log_list')

    entry_type = get_object_or_404(PrincipalLogEntryType, id=type_id, school_name=sdata)
    name = entry_type.name
    entry_type.delete()
    messages.success(request, f'Type "{name}" deleted.')

    # return to wherever the request came from
    referer = request.META.get('HTTP_REFERER', '')
    if 'types' in referer:
        return redirect('manage_principal_log_types')
    return redirect('principal_log_list')


# =========================================
# PRINCIPAL LOG ACCESS — MANAGE
# =========================================

def manage_principal_log_access(request):
    usr = request.user
    if not usr.groups.filter(name__in=['superadmin', 'Admin']).exists():
        messages.error(request, 'You do not have permission to manage access.')
        return redirect('principal_log_list')

    sch_id = request.session['sch_id']
    sdata = school.objects.get(pk=sch_id)

    staff_qs = staff.objects.filter(staff_school=sdata).order_by('first_name')

    if request.method == 'POST':
        form = PrincipalLogAccessForm(request.POST)
        form.fields['staff_members'].queryset = staff_qs
        if form.is_valid():
            selected = form.cleaned_data['staff_members']
            added, skipped = 0, 0
            for s in selected:
                _, created = PrincipalLogAccessMapping.objects.get_or_create(
                    school_name=sdata,
                    staff_member=s
                )
                if created:
                    added += 1
                else:
                    skipped += 1
            if added:
                messages.success(request, f'{added} staff member(s) granted access.')
            if skipped:
                messages.warning(request, f'{skipped} already had access and were skipped.')
        else:
            messages.error(request, 'Please correct the errors below.')
        return redirect('manage_principal_log_access')

    form = PrincipalLogAccessForm()
    form.fields['staff_members'].queryset = staff_qs

    mappings = PrincipalLogAccessMapping.objects.filter(
        school_name=sdata
    ).select_related('staff_member').order_by('staff_member__first_name')

    return render(request, 'jacobreports/principal_log_access.html', {
        'form': form,
        'mappings': mappings,
        'sdata': sdata,
    })


# =========================================
# PRINCIPAL LOG ACCESS — REVOKE
# =========================================

def revoke_principal_log_access(request, mapping_id):
    usr = request.user
    if not usr.groups.filter(name__in=['superadmin', 'Admin']).exists():
        messages.error(request, 'You do not have permission to perform this action.')
        return redirect('principal_log_list')

    mapping = get_object_or_404(PrincipalLogAccessMapping, id=mapping_id)
    mapping.delete()
    messages.success(request, 'Access revoked successfully.')
    return redirect('manage_principal_log_access')